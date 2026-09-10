import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from process_data import PRINCIPLES_DATA, POLICIES_DATA, RACI_DATA, INDEX_DATA, ALL_SECTIONS

# ---------------------------------------------------------------------------
# COLOR PALETTE CONSTANTS
# ---------------------------------------------------------------------------
COLOR_NAVY_HEX = "1B365D"       # #1B365D - Deep Navy (Header)
COLOR_SLATE_HEX = "2C3E50"      # #2C3E50 - Dark Slate (Section Banner)
COLOR_L3_HEX = "34495E"         # #34495E - Medium Slate (L3 Banner)
COLOR_LIGHT_ZEBRA_HEX = "F8F9FA"# #F8F9FA - Light Gray Zebra
COLOR_HIGHLIGHT_HEX = "EBF5FB"  # #EBF5FB - Ice Blue Highlight
COLOR_BORDER_HEX = "D5D8DC"     # #D5D8DC - Subtle Border Gray
COLOR_TEXT_DARK_HEX = "2C3E50"  # #2C3E50 - Primary Dark Text
COLOR_CODE_HEX = "1A5276"       # #1A5276 - Code / Accent Blue

RGB_NAVY = RGBColor(27, 54, 93)
RGB_SLATE = RGBColor(44, 62, 80)
RGB_MUTED = RGBColor(86, 101, 115)
RGB_DARK = RGBColor(44, 62, 80)
RGB_CODE = RGBColor(26, 82, 118)
RGB_WHITE = RGBColor(255, 255, 255)

# ---------------------------------------------------------------------------
# XML FORMATTING HELPERS
# ---------------------------------------------------------------------------
def set_cell_shading(cell, color_hex):
    """Set background color of a cell (hex string without #)."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set cell padding in dxa (1 pt = 20 dxa; 100 dxa = 5 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, color="D5D8DC", sz="4", val="single"):
    """Set subtle cell borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def make_row_header(row):
    """Mark row as repeating table header and prevent row split across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    trPr.append(OxmlElement('w:cantSplit'))

def prevent_row_split(row):
    """Prevent row split across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def set_table_col_widths(table, col_widths):
    """Set fixed width on all cells in each column."""
    for row in table.rows:
        for i, w in enumerate(col_widths):
            if i < len(row.cells):
                row.cells[i].width = Inches(w)

def add_heading_styled(doc, text, level, space_before=14, space_after=6):
    """Add a clean, styled heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = RGB_NAVY
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGB_SLATE
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGB_CODE
    return p

# ---------------------------------------------------------------------------
# DOCUMENT GENERATION MAIN
# ---------------------------------------------------------------------------
def generate_document():
    doc = docx.Document()
    
    # 1. Page Setup: Landscape Orientation, 0.5 in margins (Total printable width = 10.0 inches)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    
    # Configure Normal Style
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(9.5)
    font_normal.color.rgb = RGB_DARK
    
    # -----------------------------------------------------------------------
    # DOCUMENT HEADER / BANNER
    # -----------------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("SAP Business One Enterprise Supply Chain: Process Design Matrix")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(18)
    run_title.bold = True
    run_title.font.color.rgb = RGB_NAVY
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Master Governance Framework: 3-Tier Principles, Policies & Operational Rules, Hierarchy Register, and 28 Level 3 Route Accounting Matrix")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10.5)
    run_sub.italic = True
    run_sub.font.color.rgb = RGB_MUTED
    
    # -----------------------------------------------------------------------
    # SECTION 1: PROCESS DESIGN PRINCIPLES, POLICIES & OPERATIONAL RULES
    # -----------------------------------------------------------------------
    add_heading_styled(doc, "1. Enterprise Governance Architecture: Principles, Policies & Operational Rules", level=1, space_before=6, space_after=4)
    
    desc_p1 = doc.add_paragraph()
    desc_p1.paragraph_format.space_after = Pt(4)
    desc_p1.add_run(
        "The enterprise supply chain architecture is governed by a strict three-tiered framework that clearly separates strategic philosophy from governance mandates and deterministic execution logic:"
    )
    
    # Tier 1 Heading & Table
    add_heading_styled(doc, "1.1 Tier 1: Core Design Principles (Strategic Philosophy & Intent)", level=2, space_before=6, space_after=3)
    
    headers_principles = ["Design Principle", "Strategic Architectural Intent", "Business Purpose & Risk Mitigation", "SAP B1 System Architectural Impact"]
    widths_principles = [1.8, 2.7, 2.7, 2.8]  # Total 10.0 in
    
    table_p = doc.add_table(rows=1, cols=len(headers_principles))
    table_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_p.autofit = False
    
    hdr_p_cells = table_p.rows[0].cells
    make_row_header(table_p.rows[0])
    for i, h in enumerate(headers_principles):
        hdr_p_cells[i].text = h
        set_cell_shading(hdr_p_cells[i], COLOR_NAVY_HEX)
        set_cell_margins(hdr_p_cells[i], top=110, bottom=110, left=130, right=130)
        set_cell_borders(hdr_p_cells[i], color=COLOR_NAVY_HEX)
        p = hdr_p_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.0)
            r.bold = True
            r.font.color.rgb = RGB_WHITE
            
    for row_idx, data in enumerate(PRINCIPLES_DATA):
        row = table_p.add_row()
        prevent_row_split(row)
        cells = row.cells
        bg_color = COLOR_LIGHT_ZEBRA_HEX if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(data):
            cells[i].text = val
            set_cell_shading(cells[i], bg_color)
            set_cell_margins(cells[i], top=80, bottom=80, left=120, right=120)
            set_cell_borders(cells[i], color=COLOR_BORDER_HEX)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGB_CODE
                else:
                    r.font.color.rgb = RGB_DARK
                    
    set_table_col_widths(table_p, widths_principles)
    
    # Tier 2 & 3 Heading & Table
    add_heading_styled(doc, "1.2 Tier 2 & 3: Master Policies & Operational Execution Rules Matrix", level=2, space_before=10, space_after=3)
    
    headers_sec1 = ["Governance Code", "Policy / Rule Name", "Scope & Applicability", "Mandatory Policy Statement & Operational Rules", "SAP B1 System Implementation & G/L Impact"]
    widths_sec1 = [1.1, 1.8, 1.4, 3.0, 2.7]  # Total 10.0 in
    
    table1 = doc.add_table(rows=1, cols=len(headers_sec1))
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False
    
    # Header Row
    hdr_cells1 = table1.rows[0].cells
    make_row_header(table1.rows[0])
    for i, h in enumerate(headers_sec1):
        hdr_cells1[i].text = h
        set_cell_shading(hdr_cells1[i], COLOR_NAVY_HEX)
        set_cell_margins(hdr_cells1[i], top=110, bottom=110, left=130, right=130)
        set_cell_borders(hdr_cells1[i], color=COLOR_NAVY_HEX)
        p = hdr_cells1[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.0)
            r.bold = True
            r.font.color.rgb = RGB_WHITE
            
    # Data Rows
    for row_idx, data in enumerate(POLICIES_DATA):
        row = table1.add_row()
        prevent_row_split(row)
        cells = row.cells
        bg_color = COLOR_LIGHT_ZEBRA_HEX if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(data):
            cells[i].text = val
            set_cell_shading(cells[i], bg_color)
            set_cell_margins(cells[i], top=80, bottom=80, left=120, right=120)
            set_cell_borders(cells[i], color=COLOR_BORDER_HEX)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGB_CODE
                elif i == 1:
                    r.bold = True
                    r.font.color.rgb = RGB_NAVY
                else:
                    r.font.color.rgb = RGB_DARK
                    
    set_table_col_widths(table1, widths_sec1)
    
    # 1.3 RACI Governance Matrix Heading & Table
    add_heading_styled(doc, "1.3 Tier 4: Enterprise RACI Governance Matrix (Role Accountability & Document Segregation)", level=2, space_before=10, space_after=3)
    
    headers_raci = ["Supply Chain Functional Activity", "ERP Doc", "Curr", "Accountable & Responsible Role", "RACI Status", "Consulted (C) & Informed (I)", "Operational Responsibility Scope & Transactional Boundary"]
    widths_raci = [1.8, 0.8, 0.5, 1.4, 1.4, 1.6, 2.5]  # Total 10.0 in
    
    table_raci = doc.add_table(rows=1, cols=len(headers_raci))
    table_raci.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_raci.autofit = False
    
    hdr_raci_cells = table_raci.rows[0].cells
    make_row_header(table_raci.rows[0])
    for i, h in enumerate(headers_raci):
        hdr_raci_cells[i].text = h
        set_cell_shading(hdr_raci_cells[i], COLOR_NAVY_HEX)
        set_cell_margins(hdr_raci_cells[i], top=110, bottom=110, left=130, right=130)
        set_cell_borders(hdr_raci_cells[i], color=COLOR_NAVY_HEX)
        p = hdr_raci_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.bold = True
            r.font.color.rgb = RGB_WHITE
            
    for row_idx, data in enumerate(RACI_DATA):
        row = table_raci.add_row()
        prevent_row_split(row)
        cells = row.cells
        bg_color = COLOR_LIGHT_ZEBRA_HEX if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(data):
            cells[i].text = val
            set_cell_shading(cells[i], bg_color)
            set_cell_margins(cells[i], top=70, bottom=70, left=100, right=100)
            set_cell_borders(cells[i], color=COLOR_BORDER_HEX)
            p = cells[i].paragraphs[0]
            if i in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.0)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGB_CODE
                elif i == 1:
                    r.font.name = "Consolas"
                    r.bold = True
                    r.font.color.rgb = RGB_CODE
                elif i == 3:
                    r.bold = True
                    r.font.color.rgb = RGB_NAVY
                elif i == 4:
                    r.bold = True
                    r.font.color.rgb = RGB_DARK
                else:
                    r.font.color.rgb = RGB_DARK
                    
    set_table_col_widths(table_raci, widths_raci)
    
    # -----------------------------------------------------------------------
    # SECTION 2: PROCESS HIERARCHY REGISTER (INDEX)
    # -----------------------------------------------------------------------
    doc.add_page_break()
    add_heading_styled(doc, "2. Process Hierarchy Register (Master Index of 28 Level 3 Routes)", level=1, space_before=4, space_after=4)
    
    desc_p2 = doc.add_paragraph()
    desc_p2.paragraph_format.space_after = Pt(6)
    desc_p2.add_run(
        "The master hierarchy register classifies all 28 Level 3 operating routes across 5 functional node groups. "
        "Every sales route strictly operates on Universal Direct AR Invoicing (Policy 6) without standalone ODLN documents."
    )
    
    headers_sec2 = ["Level 2 Group", "L3 Code", "Process Pathway & Title", "Operating DBs", "Origin", "Transit Node", "Destination", "Costing & Valuation Mechanism", "Key SAP B1 Documents"]
    widths_sec2 = [1.2, 0.6, 1.8, 1.0, 0.8, 1.0, 0.9, 1.6, 1.1]  # Total 10.0 in
    
    table2 = doc.add_table(rows=1, cols=len(headers_sec2))
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    
    hdr_cells2 = table2.rows[0].cells
    make_row_header(table2.rows[0])
    for i, h in enumerate(headers_sec2):
        hdr_cells2[i].text = h
        set_cell_shading(hdr_cells2[i], COLOR_NAVY_HEX)
        set_cell_margins(hdr_cells2[i], top=110, bottom=110, left=120, right=120)
        set_cell_borders(hdr_cells2[i], color=COLOR_NAVY_HEX)
        p = hdr_cells2[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.0)
            r.bold = True
            r.font.color.rgb = RGB_WHITE
            
    for row_idx, data in enumerate(INDEX_DATA):
        row = table2.add_row()
        prevent_row_split(row)
        cells = row.cells
        bg_color = COLOR_LIGHT_ZEBRA_HEX if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(data):
            cells[i].text = val
            set_cell_shading(cells[i], bg_color)
            set_cell_margins(cells[i], top=70, bottom=70, left=110, right=110)
            set_cell_borders(cells[i], color=COLOR_BORDER_HEX)
            p = cells[i].paragraphs[0]
            if i in [1, 3, 4, 5, 6]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.0)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGB_SLATE
                elif i == 1:
                    r.bold = True
                    r.font.color.rgb = RGB_CODE
                elif i == 2:
                    r.bold = True
                    r.font.color.rgb = RGB_NAVY
                else:
                    r.font.color.rgb = RGB_DARK
                    
    set_table_col_widths(table2, widths_sec2)
    
    # -----------------------------------------------------------------------
    # SECTION 3: LEVEL 4 PROCESS STEP DETAILS (ALL 28 PROCESSES)
    # -----------------------------------------------------------------------
    doc.add_page_break()
    add_heading_styled(doc, "3. Level 4 Process Step Execution & Accounting Matrix", level=1, space_before=4, space_after=4)
    
    desc_p3 = doc.add_paragraph()
    desc_p3.paragraph_format.space_after = Pt(6)
    desc_p3.add_run(
        "Complete procedural breakdown of all 28 Level 3 supply network routes and control workflows. "
        "Each milestone follows the standard naming convention: [Owner] [Verb] [Input] [Where] for [Action/What], detailing "
        "database routing, warehouse custody (OITW.OnHand), valuation formulas, and exact double-entry Chart of Accounts (CoA) G/L postings (OJDT)."
    )
    
    headers_sec3 = [
        "L3 Code", 
        "Step Code", 
        "Step Name & Milestone [Owner + Verb + Input + Where + Action]", 
        "SAP B1 Doc", 
        "Database", 
        "Whs Node", 
        "OnHand Snapshot (OITW)", 
        "1) Inventory Valuation & Data Fields", 
        "2) CoA G/L Journal Entry (OJDT)"
    ]
    widths_sec3 = [0.55, 0.65, 2.0, 0.9, 0.8, 0.9, 1.0, 1.6, 1.6]  # Total 10.0 in
    
    for sec in ALL_SECTIONS:
        # Section Heading
        add_heading_styled(doc, sec["section_banner"], level=2, space_before=14, space_after=6)
        
        for proc in sec["processes"]:
            # Process Subheading
            add_heading_styled(doc, proc["l3_banner"], level=3, space_before=10, space_after=4)
            
            # Create Step Table for this process
            table3 = doc.add_table(rows=1, cols=len(headers_sec3))
            table3.alignment = WD_TABLE_ALIGNMENT.CENTER
            table3.autofit = False
            
            # Header Row
            hdr_cells3 = table3.rows[0].cells
            make_row_header(table3.rows[0])
            for i, h in enumerate(headers_sec3):
                hdr_cells3[i].text = h
                set_cell_shading(hdr_cells3[i], COLOR_NAVY_HEX)
                set_cell_margins(hdr_cells3[i], top=100, bottom=100, left=90, right=90)
                set_cell_borders(hdr_cells3[i], color=COLOR_NAVY_HEX)
                p = hdr_cells3[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(8.0)
                    r.bold = True
                    r.font.color.rgb = RGB_WHITE
                    
            for row_idx, step_data in enumerate(proc["steps"]):
                row = table3.add_row()
                prevent_row_split(row)
                cells = row.cells
                bg_color = COLOR_LIGHT_ZEBRA_HEX if row_idx % 2 == 1 else "FFFFFF"
                
                for i, val in enumerate(step_data):
                    cells[i].text = str(val)
                    set_cell_shading(cells[i], bg_color)
                    set_cell_margins(cells[i], top=70, bottom=70, left=80, right=80)
                    set_cell_borders(cells[i], color=COLOR_BORDER_HEX)
                    p = cells[i].paragraphs[0]
                    if i in [0, 1]:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif i in [3, 4, 5]:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        
                    for r in p.runs:
                        r.font.name = "Calibri"
                        r.font.size = Pt(7.5)
                        if i in [0, 1]:
                            r.bold = True
                            r.font.color.rgb = RGB_CODE
                        elif i in [2]:
                            r.bold = True
                            r.font.color.rgb = RGB_NAVY
                        elif i in [3, 4, 5]:
                            r.bold = True
                            r.font.color.rgb = RGB_SLATE
                        elif i == 6:
                            r.font.color.rgb = RGB_MUTED
                        else:
                            r.font.color.rgb = RGB_DARK
                            
            set_table_col_widths(table3, widths_sec3)
            
            # Spacer after table
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(6)
            
    output_docx = "/Users/billy/SAPB1_Sim/SAP_B1_Process_Design_Matrix.docx"
    doc.save(output_docx)
    print(f"Successfully generated Word document at: {output_docx}")

if __name__ == "__main__":
    generate_document()
